# md (tiny resume dialect) -> tight 2-page .docx. Arial; Liberation Sans is metric-identical, so local render = Word layout.
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = sys.argv[1], sys.argv[2]
BODY = float(sys.argv[3]) if len(sys.argv) > 3 else 10.5
LEAD = float(sys.argv[4]) if len(sys.argv) > 4 else 1.1   # line spacing multiple
MARG = float(sys.argv[5]) if len(sys.argv) > 5 else 0.65  # side margins, inches
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin"):
    setattr(s, side, Inches(MARG))
s.top_margin = s.bottom_margin = Inches(0.55)

st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(BODY)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(0)
from docx.enum.text import WD_LINE_SPACING
st.paragraph_format.line_spacing = Pt(round(BODY * LEAD * 1.15 * 2) / 2)  # LEAD x Arial's ~1.15em single line, to 0.5pt
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY

NBH = "\u2011"  # write U+2011 in the md; emitted as Word's native w:noBreakHyphen (glyph-safe in Arial)
def nbh_run(p, text):
    r = p.add_run()
    for i, part in enumerate(text.split(NBH)):
        if i: r._r.append(OxmlElement("w:noBreakHyphen"))
        if part:
            t = OxmlElement("w:t"); t.text = part; t.set(qn("xml:space"), "preserve"); r._r.append(t)
    return r

def runs(p, text, base_bold=False, base_italic=False):
    # **bold**, *italic*, [txt](url) -> visible text only (ATS-safe, no hyperlink field codes)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok: continue
        if tok.startswith("**"): r = nbh_run(p, tok[2:-2]); r.bold = True
        elif tok.startswith("*"): r = nbh_run(p, tok[1:-1]); r.italic = True
        else: r = nbh_run(p, tok)
        if base_bold: r.bold = True
        if base_italic: r.italic = True
    return p

def rule_below(p):
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3A5F")): b.set(qn("w:" + k), v)
    bdr.append(b); pPr.append(bdr)

RIGHT = Inches(8.5 - 2 * MARG)  # right tab stop = text width
def two_col(left, right, bold=False, italic=False, before=0):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT, WD_TAB_ALIGNMENT.RIGHT)
    runs(p, left, bold, italic)
    if right: p.add_run("\t" + right)
    p.paragraph_format.keep_with_next = True

for line in open(SRC, encoding="utf-8").read().splitlines():
    if not line.strip(): continue
    if line.startswith("# "):
        p = doc.add_paragraph(); r = p.add_run(line[2:]); r.bold = True; r.font.size = Pt(BODY + 8); r.font.color.rgb = ACCENT
    elif line.startswith("> "):
        p = doc.add_paragraph(); r = p.add_run(line[2:]); r.font.size = Pt(BODY + 2); r.font.color.rgb = ACCENT
        p.paragraph_format.space_after = Pt(2)
    elif line.startswith("~ "):
        p = runs(doc.add_paragraph(), line[2:]); pf = p.paragraph_format
        pf.left_indent = Inches(0.18); pf.space_after = Pt(2.5)
    elif line.startswith("## "):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
        r = p.add_run(line[3:].upper()); r.bold = True; r.font.size = Pt(BODY + 1); r.font.color.rgb = ACCENT
        rule_below(p); p.paragraph_format.keep_with_next = True
    elif line.startswith("### "):
        brk = line[4:].startswith("!"); l, _, rt = line[4:].lstrip("!").partition(" | "); two_col(l, rt, bold=True, before=7)
        if brk: doc.paragraphs[-1].paragraph_format.page_break_before = True
    elif line.startswith("= "):
        l, _, rt = line[2:].partition(" | "); two_col(l, rt)
        pf = doc.paragraphs[-1].paragraph_format; pf.space_after = Pt(2); pf.keep_with_next = False
    elif line.startswith("@ "):
        l, _, rt = line[2:].partition(" | "); two_col(l, rt, italic=True)
    elif line.startswith("- "):
        p = runs(doc.add_paragraph(style="List Bullet"), line[2:])
        pf = p.paragraph_format; pf.left_indent = Inches(0.18); pf.first_line_indent = Inches(-0.13); pf.space_after = Pt(2.5)
    else:
        p = runs(doc.add_paragraph(), line); p.paragraph_format.space_after = Pt(2)

doc.save(OUT)
